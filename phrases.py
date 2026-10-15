#!/usr/bin/env python3
"""Extract the TC/SE signal phrases from the strategy-series workbook.

    python phrases.py "<path to the .xlsx>"

The 'All Phrases' sheet is three columns: the phrase, what it signals, and an
example sentence. What it signals is the thing being tested - whether a blank
continues the sentence's logic, reverses it, negates it, or leaves it alone -
so it takes the place of a definition, and the sentence takes the place of the
extended meaning.

Workbooks are gitignored (GregMat's paid material); phrases.json is committed.
"""
from __future__ import annotations

import json
import pathlib
import sys

import openpyxl

HERE = pathlib.Path(__file__).parent
OUT = HERE / "phrases.json"
SHEET = "All Phrases"


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    src = pathlib.Path(sys.argv[1])
    if not src.exists():
        print(f"not found: {src}")
        return 1

    wb = openpyxl.load_workbook(src, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"no {SHEET!r} sheet; found {wb.sheetnames}")
        return 1

    rows, skipped = [], 0
    for r in wb[SHEET].iter_rows(min_row=2, values_only=True):
        phrase = str(r[0]).strip() if r[0] else ""
        kind = str(r[1]).strip() if len(r) > 1 and r[1] else ""
        sentence = str(r[2]).strip() if len(r) > 2 and r[2] else ""
        if not phrase or not kind:
            skipped += 1
            continue
        rows.append({"p": phrase, "t": kind, "s": sentence})

    OUT.write_text(json.dumps(rows, ensure_ascii=False, indent=0), encoding="utf-8")
    kinds: dict[str, int] = {}
    for x in rows:
        kinds[x["t"]] = kinds.get(x["t"], 0) + 1
    print(f"{len(rows)} phrases -> {OUT}")
    for k, n in sorted(kinds.items()):
        print(f"  {k}: {n}")
    if skipped:
        print(f"  ({skipped} rows skipped for a missing phrase or type)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
